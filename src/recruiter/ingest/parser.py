"""Stage 1a — resume file to plain text.

Handles PDF, DOCX, TXT and MD. Returns warnings instead of raising so one bad
file cannot take down a batch run; the pipeline records the warning and moves on.

The scanned-PDF case matters: a PDF with no text layer parses "successfully"
into an empty string. We detect that and flag it for manual review rather than
letting an empty resume score zero and get silently rejected.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}

# Below this, the file almost certainly did not parse properly.
MIN_USABLE_CHARS = 120


class UnsupportedResumeFormat(ValueError):
    pass


@dataclass
class ParsedResume:
    path: Path
    text: str = ""
    warnings: list[str] = field(default_factory=list)

    @property
    def usable(self) -> bool:
        return len(self.text.strip()) >= MIN_USABLE_CHARS

    @property
    def filename(self) -> str:
        return self.path.name


def _clean(text: str) -> str:
    """Normalise whitespace without destroying the line structure that section
    headers and bullet points depend on."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace(" ", " ").replace("﻿", "")
    # PDF extraction often leaves ligatures and bullet glyphs
    for bad, good in (("ﬁ", "fi"), ("ﬂ", "fl"), ("•", "-"), ("–", "-")):
        text = text.replace(bad, good)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def _parse_pdf(path: Path, warnings: list[str]) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        warnings.append("pypdf not installed — cannot read PDF. Run pip install -r requirements.txt")
        return ""

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        warnings.append(f"PDF could not be opened ({exc}). File may be corrupt or encrypted.")
        return ""

    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")  # many resumes are "encrypted" with an empty password
        except Exception:
            warnings.append("PDF is password protected — needs manual review.")
            return ""

    pages: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:
            warnings.append(f"page {i} could not be extracted ({exc})")

    text = "\n".join(pages)
    if len(text.strip()) < MIN_USABLE_CHARS:
        warnings.append(
            "PDF has little or no text layer — it is probably a scan or an image. "
            "Needs OCR or manual review."
        )
    return text


def _parse_docx(path: Path, warnings: list[str]) -> str:
    try:
        import docx
    except ImportError:
        warnings.append(
            "python-docx not installed — cannot read DOCX. Run pip install -r requirements.txt"
        )
        return ""

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        warnings.append(f"DOCX could not be opened ({exc})")
        return ""

    parts = [p.text for p in document.paragraphs]
    # Plenty of resumes put the whole layout in tables; paragraphs alone miss those.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _parse_text(path: Path, warnings: list[str]) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
        except OSError as exc:
            warnings.append(f"could not read file ({exc})")
            return ""
    warnings.append("could not decode file with any known encoding")
    return ""


def parse_resume(path: Path | str) -> ParsedResume:
    """Read one resume file. Never raises for content problems — check `.warnings`."""
    path = Path(path)
    warnings: list[str] = []

    if not path.exists():
        return ParsedResume(path=path, warnings=[f"file not found: {path}"])
    if path.stat().st_size == 0:
        return ParsedResume(path=path, warnings=["file is empty (0 bytes)"])

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        return ParsedResume(
            path=path,
            warnings=[
                f"unsupported format '{suffix}'. "
                f"Supported: {', '.join(sorted(SUPPORTED_SUFFIXES))}"
            ],
        )

    if suffix == ".pdf":
        raw = _parse_pdf(path, warnings)
    elif suffix == ".docx":
        raw = _parse_docx(path, warnings)
    else:
        raw = _parse_text(path, warnings)

    text = _clean(raw)
    if text and len(text) < MIN_USABLE_CHARS and not warnings:
        warnings.append(
            f"resume is only {len(text)} characters — too short to assess fairly, "
            "flagged for manual review"
        )
    return ParsedResume(path=path, text=text, warnings=warnings)


def discover_resumes(folder: Path | str) -> list[Path]:
    """All supported resume files in a folder, sorted for reproducible runs."""
    folder = Path(folder)
    if not folder.is_dir():
        raise NotADirectoryError(f"resume folder not found: {folder}")
    return sorted(
        p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES
    )
